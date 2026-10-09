#!/usr/bin/env python3
"""받기 파일(/files/)과 아이콘·공유 그림을 다시 만든다. 사이트 화면과 같은 배치 코드(assets/sheet.js)로 그린다.

  python3 calendar/_dev/make_files.py            # PDF 27개 + 엑셀 3개 (pdf 또는 xlsx 만 주면 그것만)
  python3 calendar/_dev/make_files.py icons      # favicon.svg/.ico, apple-touch-icon.png, og.png, og-ko.png

필요한 것: playwright(크로미움), pypdf, openpyxl, Pillow, 그리고 Pretendard 가변 글꼴 파일(woff2).
PDF는 크로미움이 글자 그대로 만들고, pypdf 로 문서 정보(제목, 만든 곳 = 사이트 이름)를 적는다.
내 설정으로 받는 PDF는 화면에서 assets/pdf.js 가 따로 만든다(그 글꼴은 make_fonts.py).
글꼴 경로는 환경 변수 CAL_FONT (없으면 /home/claude/sp/calendar/mock/shared/fonts/PretendardVariable.woff2).
받는 파일 안에는 광고가 없다. 구석에 사이트 주소 한 줄만 아주 작게 들어간다.
"""
import sys
sys.dont_write_bytecode = True
import json
import os
import subprocess
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEV = ROOT / '_dev'
FONT = os.environ.get('CAL_FONT', '/home/claude/sp/calendar/mock/shared/fonts/PretendardVariable.woff2')
CHROME = os.environ.get('CAL_CHROME', '/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
YEARS = [2026, 2027, 2028]
BOTH = {'1월 1일': '신정', '3·1절': '삼일절', '기독탄신일': '성탄절'}   # 규정의 이름 → 흔히 쓰는 이름(build.py KR_BOTH 와 같다)
FACE = '@font-face{font-family:"Pretendard Variable";font-weight:45 920;font-style:normal;src:url("file://%s") format("woff2-variations")}' % FONT


def node(jobs):
    return json.loads(subprocess.run(['node', str(DEV / 'sheets.js')], input=json.dumps(jobs), capture_output=True, text=True, check=True).stdout)


def jobs():
    out = []
    for y in YEARS:
        for lang, c, papers in (('ko', 'KR', ['a4']), ('en', 'US', ['letter', 'a4'])):
            cname = 'korea' if c == 'KR' else 'us'
            for paper in papers:
                base = dict(year=y, lang=lang, country=c, paper=paper)
                for orient in ('landscape', 'portrait'):
                    out.append(dict(name=f'{y}-calendar-{cname}-{paper}-{orient}.pdf', lang=lang, pages=[dict(base, kind='year', orient=orient)]))
                show = dict(names=True, lunar=True, terms=True) if c == 'KR' else dict(names=True)
                out.append(dict(name=f'{y}-calendar-{cname}-monthly-{paper}.pdf', lang=lang, title=(f'{y}년 월별 달력' if lang == 'ko' else f'{y} monthly calendar'),
                                pages=[dict(base, kind='month', month=m, show=show) for m in range(1, 13)]))
    return out


def make_pdfs():
    from playwright.sync_api import sync_playwright
    from pypdf import PdfReader, PdfWriter
    files = ROOT / 'files'
    files.mkdir(exist_ok=True)
    todo = jobs()
    res = node(todo)
    for r, j in zip(res, todo):
        r['lang'] = j['lang']
    with tempfile.TemporaryDirectory() as tmp, sync_playwright() as p:
        b = p.chromium.launch(executable_path=CHROME)
        pg = b.new_page()
        for r in res:
            w, h = r['w'], r['h']
            html = ('<!doctype html><html><head><meta charset="utf-8"><title>%s</title><style>%s@page{size:%smm %smm;margin:0}html,body{margin:0}'
                    'svg{display:block;width:%smm;height:%smm;break-after:page}svg:last-child{break-after:auto}</style></head><body>%s</body></html>') % (
                r['title'], FACE, w, h, w, h - 0.3, ''.join(r['svgs']))
            f = Path(tmp) / 'p.html'
            f.write_text(html, encoding='utf-8')
            pg.goto('file://' + str(f))
            pg.evaluate('document.fonts.ready')
            raw = Path(tmp) / 'p.pdf'
            pg.pdf(path=str(raw), prefer_css_page_size=True, print_background=True)
            # 문서 정보: 제목과 만든 곳(사이트 이름). 크로미움이 적은 'HeadlessChrome'·'Skia/PDF'는 남기지 않는다.
            who = ('한장달력' if r['lang'] == 'ko' else 'Onesheet') + ' (calendar.lumenlab.page)'
            w = PdfWriter(clone_from=PdfReader(str(raw)))
            w.add_metadata({'/Title': r['title'], '/Author': who, '/Creator': who, '/Producer': who})
            w.write(str(files / r['name']))
            print(r['name'], (files / r['name']).stat().st_size, len(w.pages))
        b.close()


def make_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    files = ROOT / 'files'
    files.mkdir(exist_ok=True)
    for data in node([dict(xlsx=y) for y in YEARS]):
        y = data['xlsx']
        wb = Workbook()
        ws0 = wb.active
        ws0.title = '공휴일'
        ws0.append([f'{y}년 공휴일' + (' (공식 발표 전 자료)' if data['basis'] == 'provisional' else '')])
        ws0['A1'].font = Font(size=16, bold=True)
        ws0.append(['날짜', '요일', '이름'])
        for c in ws0[2]:
            c.font = Font(bold=True)
        for it in data['list']:   # 이름: 규정의 이름, 괄호 안에 흔히 쓰는 이름
            ws0.append([f'{it["m"]}월 {it["d"]}일', '일월화수목금토'[it['wd']], '; '.join(n + (f'({BOTH[n]})' if n in BOTH else '') for n in it['names'])])
        ws0.append([])
        for note in ('이름은 규정의 이름이고, 괄호 안은 흔히 쓰는 이름이에요. 달력 시트의 칸에는 흔히 쓰는 이름을 적었어요.',
                     '공휴일 기준: 관공서의 공휴일에 관한 규정(2026년 4월 개정 반영). 선거일·임시공휴일은 확정된 날만 들어 있어요.',
                     '일요일 시작 기본판이에요. 주 시작 요일이나 용지를 바꾼 달력은 calendar.lumenlab.page 에서 PDF로 받을 수 있어요.'):
            ws0.append([note])
            r = ws0.max_row
            ws0.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            ws0.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
            ws0.cell(row=r, column=1).font = Font(size=10, color='5F6B65')
            ws0.row_dimensions[r].height = 30
        ws0.column_dimensions['A'].width = 14
        ws0.column_dimensions['B'].width = 8
        ws0.column_dimensions['C'].width = 44
        ws0.page_setup.paperSize = ws0.PAPERSIZE_A4
        ws0.page_setup.fitToWidth = 1
        ws0.page_setup.fitToHeight = 1
        ws0.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        thin = Side(style='thin', color='C9CFCB')
        red, blue, gray = 'CF3027', '2355C8', '727B75'
        for mo in data['months']:
            ws = wb.create_sheet(f'{mo["m"]}월')
            ws.merge_cells('A1:G1')
            ws['A1'] = f'{y}년 {mo["m"]}월'
            ws['A1'].font = Font(size=20, bold=True)
            ws.row_dimensions[1].height = 34
            for i, n in enumerate('일월화수목금토'):
                c = ws.cell(row=2, column=i + 1, value=n)
                c.font = Font(bold=True, color=red if i == 0 else blue if i == 6 else '191C1A')
                c.fill = PatternFill('solid', fgColor='EEF2EF')
                c.alignment = Alignment(horizontal='center')
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            days = {d['d']: d for d in mo['days']}
            for r, row in enumerate(mo['grid']):
                ws.row_dimensions[3 + r].height = 66
                for i, d in enumerate(row):
                    c = ws.cell(row=3 + r, column=i + 1)
                    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    c.alignment = Alignment(vertical='top', wrap_text=True)
                    if not d:
                        continue
                    info = days[d]
                    extra = [x for x in (info['hol'], info['term'], info['lunar']) if x]
                    c.value = str(d) + ('\n' + '\n'.join(extra) if extra else '')
                    c.font = Font(size=12, bold=bool(info['hol']), color=red if (i == 0 or info['hol']) else blue if i == 6 else '191C1A')
            for i in range(7):
                ws.column_dimensions['ABCDEFG'[i]].width = 17
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        name = files / f'{y}-calendar-korea.xlsx'
        wb.save(name)
        print(name.name, name.stat().st_size)


ICON = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 32 32"><path d="M6 2.5h13.5L27 10v19.5H6z" fill="#fff" stroke="#17201b" stroke-width="2.2" stroke-linejoin="round"/>'
        '<path d="M19.5 2.5V10H27z" fill="#0e6a51" stroke="#17201b" stroke-width="2.2" stroke-linejoin="round"/><g fill="#17201b"><rect x="9.6" y="14" width="3.2" height="3.2" rx=".7"/>'
        '<rect x="15" y="14" width="3.2" height="3.2" rx=".7"/><rect x="9.6" y="19.2" width="3.2" height="3.2" rx=".7"/><rect x="15" y="19.2" width="3.2" height="3.2" rx=".7"/>'
        '<rect x="20.4" y="19.2" width="3.2" height="3.2" rx=".7"/><rect x="9.6" y="24.4" width="3.2" height="3.2" rx=".7"/></g><rect x="20.4" y="14" width="3.2" height="3.2" rx=".7" fill="#cf3027"/></svg>')


def make_icons():
    from PIL import Image
    from playwright.sync_api import sync_playwright
    (ROOT / 'favicon.svg').write_text(ICON, encoding='utf-8')
    sheets = {r['name']: r for r in node([dict(name='ko', pages=[dict(kind='year', year=2027, lang='ko')]), dict(name='en', pages=[dict(kind='year', year=2027, lang='en', paper='letter')])])}
    with tempfile.TemporaryDirectory() as tmp, sync_playwright() as p:
        b = p.chromium.launch(executable_path=CHROME)
        pg = b.new_page(viewport={'width': 512, 'height': 512})
        pg.set_content(f'<body style="margin:0;background:transparent"><div style="width:512px;height:512px">{ICON}</div></body>')
        pg.screenshot(path=f'{tmp}/i.png', omit_background=True)
        im = Image.open(f'{tmp}/i.png').convert('RGBA')
        im.resize((32, 32), Image.LANCZOS).save(ROOT / 'favicon.ico', sizes=[(32, 32), (16, 16)])
        pg.set_content(f'<body style="margin:0;background:#f1f4f0"><div style="width:512px;height:512px;padding:66px;box-sizing:border-box">{ICON}</div></body>')
        pg.screenshot(path=f'{tmp}/a.png')
        Image.open(f'{tmp}/a.png').convert('RGB').resize((180, 180), Image.LANCZOS).save(ROOT / 'apple-touch-icon.png', optimize=True)
        pg.set_viewport_size({'width': 1200, 'height': 630})
        for lang, out, h1, sub, brand in (('ko', 'og-ko.png', '2027년 달력 인쇄', '공휴일 넣은 한 장 PDF · 가입 없이 바로', '한장달력'),
                                          ('en', 'og.png', 'Printable 2027 Calendar', 'One page with holidays · free PDF', 'Onesheet')):
            html = f'''<!doctype html><meta charset="utf-8"><style>{FACE}
body{{margin:0;width:1200px;height:630px;background:#f1f4f0;font-family:"Pretendard Variable";color:#17201b;position:relative;overflow:hidden}}
.b{{position:absolute;left:64px;top:60px;display:flex;align-items:center;gap:14px;font-size:34px;font-weight:800;letter-spacing:-.02em}}
.b svg{{width:40px;height:40px}}
h1{{position:absolute;left:64px;top:236px;margin:0;font-size:{74 if lang == 'ko' else 62}px;font-weight:800;letter-spacing:-.04em;line-height:1.1;width:520px}}
p{{position:absolute;left:66px;top:{420 if lang == 'ko' else 438}px;margin:0;font-size:28px;color:#39453f;font-weight:500;width:500px;line-height:1.35}}
.s{{position:absolute;left:600px;top:84px;width:720px;box-shadow:0 2px 2px rgba(16,32,24,.08),0 30px 60px -24px rgba(16,40,28,.45);transform:rotate(-3deg)}}
.s svg{{display:block;width:100%;height:auto}}</style>
<div class="b">{ICON}<span>{brand}</span></div><h1>{h1}</h1><p>{sub}</p><div class="s">{sheets[lang]['svgs'][0]}</div>'''
            f = Path(tmp) / 'og.html'
            f.write_text(html, encoding='utf-8')
            pg.goto('file://' + str(f))
            pg.evaluate('document.fonts.ready')
            pg.wait_for_timeout(200)
            pg.screenshot(path=str(ROOT / out))
            Image.open(ROOT / out).convert('RGB').save(ROOT / out, optimize=True)
            print(out, (ROOT / out).stat().st_size)
        b.close()


if __name__ == '__main__':
    if not Path(FONT).exists():
        sys.exit('글꼴 파일이 없어요: ' + FONT + ' (CAL_FONT 로 Pretendard 가변 woff2 경로를 주세요)')
    if 'icons' in sys.argv:
        make_icons()
    elif 'xlsx' in sys.argv:
        make_xlsx()
    elif 'pdf' in sys.argv:
        make_pdfs()
    else:
        make_pdfs()
        make_xlsx()
